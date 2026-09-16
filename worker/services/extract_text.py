"""Extract text from PDF and DOCX files."""

from pathlib import Path


def extract_text(file_path: Path) -> str:
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        return _extract_pdf(file_path)
    elif suffix in (".docx", ".doc"):
        return _extract_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {suffix}")


def _extract_pdf(path: Path) -> str:
    import pymupdf
    doc = pymupdf.open(str(path))
    parts = []
    for page in doc:
        text = page.get_text()
        if text.strip():
            parts.append(text)
    return "\n\n".join(parts)


def _extract_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Also extract tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)
    return "\n\n".join(paragraphs)
